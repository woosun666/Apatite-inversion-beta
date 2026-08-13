#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""
Ok Tedi apatite Cl-OH-F inversion - Tier 2 (follows the Lormand original workflow)
=================================================================================
Port of the 'Explore/CreateTargets' + 'Run multistart to a target' paths of
Apatite_V1_SNCC_OR_26214.m. Same TWO-STAGE design as the original (NOT a direct
scatter fit):

  1) create-target : hand-set TARGET_PARAMS -> forward model -> a TARGET CURVE
                     (Frac_Melt, XCl/XOH, XF/XOH, XF/XCl). You hand-tune the params
                     until the curve passes through the observed apatite ratios.
                     (As in the original, the observed points are overplotted for
                     the eye; they are NOT in the automatic objective.)
  2) multistart    : fit that target curve -> a family of best-fit parameter sets
                     (BOUNDS-constrained), written in the MATLAB Multioutput format
                     so the rows can be fed straight back to
                     `apatite_model.py --params`.

IMPORTED FROM SEPARATE FILES (supplied later -- see the USER-EDITABLE block):
  * OBSERVED_RATIOS_FILE : the apatite Cl-OH-F mole ratios (NOT computed here).
  * GROUPS_FILE          : ZHAI / matrix assignment per analysis (NOT decided here).

Forward physics is reused from apatite_model.py (single source of truth). The
inversion uses a fast vectorised fixed-point partition solve, cross-checked
against apatite_model's least_squares solve at startup (must agree to ~1e-9).

Deps: numpy + scipy (+ stdlib csv). xlsx import additionally needs openpyxl.
"""
from __future__ import annotations

import argparse
import csv
import os

import numpy as np
from scipy.optimize import minimize

import apatite_model as am

# =====================  USER-EDITABLE BLOCKS (edit freely)  ===================
# (1) Target forward-run parameters  == MATLAB CreateTargets InputParam.
#     Hand-tune these until the target curve matches the observed ratios.
TARGET_PARAMS = dict(
    H2O_init=4.0, Cl_init=0.24, F_init=0.14,
    D_xl_m_Cl=0.2, D_xl_m_F=0.6, D_xl_m_H2O=0.12,
    D_fl_m_Cl=20.0, D_fl_m_F=4.0, H2O_Sat=8.5,
    D_fl_m_PrimBoilH2O=0.0,          # primary boiling off (kept constant in fitting)
)
K_SPECIATION    = 0.361
NUM_FRAC_TARGET = 200                # target-curve resolution (MATLAB CreateTargets)
NUM_FRAC_FIT    = 100                # forward resolution during fitting (MATLAB)

# (2) Multistart parameter bounds  == MATLAB lb/ub (same 9 params, PARAM_ORDER below).
BOUNDS = dict(
    H2O_init=(3.7, 4.7), Cl_init=(0.16, 0.30), F_init=(0.04, 0.22),
    D_xl_m_Cl=(0.2, 0.4), D_xl_m_F=(0.5, 0.9), D_xl_m_H2O=(0.12, 0.25),
    D_fl_m_Cl=(17.0, 30.0), D_fl_m_F=(1.5, 5.0), H2O_Sat=(7.5, 10.0),
)
NUM_START = 50                       # MATLAB used 2000; raise once you trust the setup
SEED      = 0

# (3) Inputs supplied LATER (set the paths when the files exist; None = skip).
OBSERVED_RATIOS_FILE = None          # apatite ratios: csv/xlsx w/ Cl_OH, F_OH, F_Cl (+opt group)
GROUPS_FILE          = None          # ZHAI/matrix per analysis: csv/xlsx (id col + group col)
TF_CURVE_FILE        = None          # T-F curve csv (Var1=T_C, Var2=F); else CONST_T_C
CONST_T_C            = 800.0         # constant melt T if no TF curve (Ti-in-zircon, P1)
# =============================================================================

PARAM_ORDER = ["H2O_init", "Cl_init", "F_init", "D_xl_m_Cl", "D_xl_m_F",
               "D_xl_m_H2O", "D_fl_m_Cl", "D_fl_m_F", "H2O_Sat"]
# MATLAB Multioutput column names, aligned 1:1 with PARAM_ORDER (+ RMSE).
MO_COLS = ["H2O", "Cl", "F", "Dxtl-mCl", "Dxtl-mF", "Dxtl-mH2O",
           "Dfl-mCl", "Dfl-mF", "H2OSat", "RMSE"]

HERE = os.path.dirname(os.path.abspath(__file__))


# ----------------------------------------------------------------------------
# Fast vectorised forward (fixed-point partition) -> apatite ratios
# ----------------------------------------------------------------------------
def _partition_fixedpoint(XF_m, XCl_m, XOH_m, T_K, iters=1000, tol=1e-14):
    """Vectorised fixed-point solve of the non-ideal Cl-OH-F partitioning for all
    melt-fraction steps at once. Same equations/constants as
    apatite_model.solve_partition_point, iterated to the fixed point instead of
    per-point least_squares (much faster for repeated inversion forwards).
    dG coefficients + Margules W are imported from apatite_model (single source).
    Returns (XClap, XOHap, XFap, converged): converged=False means the iteration
    budget ran out before reaching tol (possible for extreme multistart draws --
    the plain fixed point is undamped and can cycle); callers in the objective
    must then reject the point rather than trust the unconverged values."""
    dG_ClOH, dG_FOH = am.delta_G(T_K)
    n = len(XF_m)
    XClap = np.full(n, 1.0 / 3.0)
    XOHap = np.full(n, 1.0 / 3.0)
    XFap = np.full(n, 1.0 / 3.0)
    converged = False
    for _ in range(iters):
        k_oh_cl = np.exp((1000.0 * (-dG_ClOH - ((XClap - XOHap) * am.WG_CLOH + XFap * am.WG_DIFF1)))
                         / (am.R_GAS * T_K))
        k_oh_f = np.exp(1000.0 * (-dG_FOH - ((XFap - XOHap) * am.WG_FOH + XClap * am.WG_DIFF2))
                        / (am.R_GAS * T_K))
        K_ClF = k_oh_f / k_oh_cl
        K_FOH = 1.0 / k_oh_f
        K_ClOH = 1.0 / k_oh_cl
        nXFap = 1.0 / (1.0 + XCl_m * K_ClF / XF_m + XOH_m / (XF_m * K_FOH))
        nXClap = 1.0 / (1.0 + XF_m / (XCl_m * K_ClF) + XOH_m / (XCl_m * K_ClOH))
        nXOHap = 1.0 / (1.0 + XF_m * K_FOH / XOH_m + XCl_m * K_ClOH / XOH_m)
        d = max(np.max(np.abs(nXFap - XFap)),
                np.max(np.abs(nXClap - XClap)),
                np.max(np.abs(nXOHap - XOHap)))
        XFap, XClap, XOHap = nXFap, nXClap, nXOHap
        if d < tol:
            converged = True
            break
    return XClap, XOHap, XFap, converged


def coeff_to_params(coeff, prim_boil=None):
    if prim_boil is None:
        prim_boil = TARGET_PARAMS["D_fl_m_PrimBoilH2O"]
    p = {k: float(coeff[i]) for i, k in enumerate(PARAM_ORDER)}
    p["D_fl_m_PrimBoilH2O"] = prim_boil
    return p


def forward_ratios(params, F, T_K, k_spec=None, with_conv=False):
    """Vectorised forward -> (XCl_div_XOH(F), XF_div_XOH(F)).
    with_conv=True additionally returns the fixed-point convergence flag
    (default off to keep the historical 2-tuple interface for the ot_test
    wrappers)."""
    if k_spec is None:
        k_spec = K_SPECIATION
    vol = am.volatile_evolution(F, params)
    _, XCl_m, XF_m, XOH_m = am.melt_mole_fractions(
        vol["H2O_rem"], vol["Cl_rem"], vol["F_rem"], k_spec)
    XClap, XOHap, XFap, conv = _partition_fixedpoint(XF_m, XCl_m, XOH_m, T_K)
    if with_conv:
        return XClap / XOHap, XFap / XOHap, conv
    return XClap / XOHap, XFap / XOHap


# ----------------------------------------------------------------------------
# Objective + multistart (Lormand: fit the TARGET CURVE, not the raw scatter)
# ----------------------------------------------------------------------------
def rmse_to_target(coeff, F_fit, T_K_fit, F_obs, tgt_cloh, tgt_foh,
                   k_spec=None, prim_boil=None, metric="rmse"):
    """Misfit of model vs target curve over XCl/XOH and XF/XOH (MATLAB line ~845).
    metric = "rmse"    : plain RMSE on the ratio values (engine default — all
                         runs up to 26.6.25 canonical used this);
             "logrmse" : RMSE on log10-transformed ratios (P1 spec §7.4 / SNCC
                         protocol wording) — down-weights the high-Cl/OH end,
                         equal RELATIVE misfit per decade.
    The two are NOT interchangeable: results/exports must state which was used.
    An unconverged partition solve returns a large penalty (1e6) so L-BFGS-B
    steps away instead of fitting to garbage."""
    cloh, foh, conv = forward_ratios(coeff_to_params(coeff, prim_boil),
                                     F_fit, T_K_fit, k_spec, with_conv=True)
    if not conv:
        return 1e6
    asc = np.argsort(F_fit)
    cloh_i = np.interp(F_obs, F_fit[asc], cloh[asc])
    foh_i = np.interp(F_obs, F_fit[asc], foh[asc])
    if metric == "logrmse":
        eps = 1e-12                     # ratios are positive; guard exact zeros
        if np.any(cloh_i <= 0) or np.any(foh_i <= 0):
            return 1e6
        res = np.concatenate([np.log10(np.maximum(tgt_cloh, eps)) - np.log10(cloh_i),
                              np.log10(np.maximum(tgt_foh, eps)) - np.log10(foh_i)])
    else:
        res = np.concatenate([tgt_cloh - cloh_i, tgt_foh - foh_i])
    return float(np.sqrt(np.mean(res ** 2)))


# --- multistart: embarrassingly parallel across start points ----------------
# Each start is an independent L-BFGS-B local solve -> fan out over processes
# (true parallelism, GIL-free). Read-only fit data is shared once per worker via
# the initializer (no per-task re-pickling); the parent draws ALL start points
# up front from the seeded RNG, so output is identical to the serial path for
# any worker count (deterministic in SEED). NB Windows uses 'spawn': workers
# re-import this module, so the worker must NOT rely on parent runtime overrides
# of the module globals -- every fit input (bounds, k_spec, prim_boil) is passed in.
_LSOPTS = dict(maxiter=300, ftol=1e-10)
_W = {}                              # per-worker shared read-only fit data


def _worker_init(F_fit, T_K_fit, F_obs, tgt_cloh, tgt_foh, bounds_list, k_spec,
                 prim_boil, lsopts=None, metric="rmse"):
    # lsopts: optional L-BFGS-B options override (e.g. loose ftol=1e-2 for a
    # Lormand-MATLAB-style exploratory fan); None -> the tight _LSOPTS default.
    # metric: objective variant ("rmse" | "logrmse"), see rmse_to_target.
    _W.update(F_fit=F_fit, T_K_fit=T_K_fit, F_obs=F_obs, tgt_cloh=tgt_cloh,
              tgt_foh=tgt_foh, bounds_list=bounds_list, k_spec=k_spec,
              prim_boil=prim_boil, lsopts=lsopts, metric=metric)


def _fit_one(x0):
    res = minimize(rmse_to_target, x0,
                   args=(_W["F_fit"], _W["T_K_fit"], _W["F_obs"], _W["tgt_cloh"],
                         _W["tgt_foh"], _W["k_spec"], _W["prim_boil"],
                         _W.get("metric", "rmse")),
                   method="L-BFGS-B", bounds=_W["bounds_list"],
                   options=_W.get("lsopts") or _LSOPTS)
    return float(res.fun), res.x


def run_multistart(F_obs, tgt_cloh, tgt_foh, F_fit, T_K_fit, num_start=None,
                   n_jobs=None, k_spec=None, prim_boil=None, ftol=None,
                   metric="rmse"):
    """Multistart fit of the target curve.
    n_jobs: None -> all cores; 1 -> serial (no subprocess); N -> N processes.
    ftol: None -> engine default (1e-10, fully converged family); a loose value
    (e.g. 1e-2 = the Lormand MATLAB GUI 'Tolerance') stops runs early so they
    fan out over the data -- DISPLAY/EXPLORATION only, not honest uncertainty.
    Deterministic in SEED (all start points drawn in the parent). Returns a list
    of (RMSE, params) sorted by ascending RMSE."""
    lb = np.array([BOUNDS[k][0] for k in PARAM_ORDER])
    ub = np.array([BOUNDS[k][1] for k in PARAM_ORDER])
    bounds_list = list(zip(lb, ub))
    ns = NUM_START if num_start is None else num_start
    if k_spec is None:
        k_spec = K_SPECIATION
    if prim_boil is None:
        prim_boil = TARGET_PARAMS["D_fl_m_PrimBoilH2O"]
    rng = np.random.default_rng(SEED)
    starts = [lb + rng.random(len(lb)) * (ub - lb) for _ in range(ns)]

    if n_jobs is None:
        n_jobs = os.cpu_count() or 1
    n_jobs = max(1, min(n_jobs, ns))
    lsopts = None if ftol is None else dict(_LSOPTS, ftol=ftol)
    init_args = (F_fit, T_K_fit, F_obs, tgt_cloh, tgt_foh, bounds_list, k_spec,
                 prim_boil, lsopts, metric)

    if n_jobs == 1:
        _worker_init(*init_args)                 # populate _W in this process
        results = [_fit_one(x0) for x0 in starts]
    else:
        from concurrent.futures import ProcessPoolExecutor
        with ProcessPoolExecutor(max_workers=n_jobs, initializer=_worker_init,
                                 initargs=init_args) as ex:
            results = list(ex.map(_fit_one, starts))
    results.sort(key=lambda t: t[0])
    return results


# ----------------------------------------------------------------------------
# Import interfaces for files supplied LATER (placeholders, but functional)
# ----------------------------------------------------------------------------
def _read_table(path):
    ext = os.path.splitext(path)[1].lower()
    if ext in (".csv", ".txt"):
        with open(path, newline="", encoding="utf-8-sig") as fh:
            return list(csv.DictReader(fh))
    if ext in (".xlsx", ".xlsm"):
        try:
            import openpyxl
        except ImportError:
            raise SystemExit("xlsx import needs openpyxl: pip install openpyxl")
        wb = openpyxl.load_workbook(path, read_only=True, data_only=True)
        rows = list(wb.active.iter_rows(values_only=True))
        hdr = [str(h) if h is not None else "" for h in rows[0]]
        return [dict(zip(hdr, r)) for r in rows[1:]]
    raise SystemExit("unknown table type: " + ext)


def _pick(row, *names):
    for n in names:
        for k, v in row.items():
            if k and str(k).strip().lower() == n.lower():
                return v
    return None


def _f(v):
    try:
        return float(v)
    except (TypeError, ValueError):
        return np.nan


def load_observed_ratios(path):
    """IMPORT observed apatite ratios from a separate file (supplied later).
    Expected columns (case-insensitive): Cl_OH, F_OH, F_Cl; optional: group, id.
    Returns dict of arrays/lists, or None if no file yet."""
    if not path or not os.path.exists(path):
        return None
    rows = _read_table(path)
    return dict(
        Cl_OH=np.array([_f(_pick(r, "Cl_OH", "XCl_div_XOH", "Cl/OH")) for r in rows]),
        F_OH=np.array([_f(_pick(r, "F_OH", "XF_div_XOH", "F/OH")) for r in rows]),
        F_Cl=np.array([_f(_pick(r, "F_Cl", "XF_div_XCl", "F/Cl")) for r in rows]),
        group=[_pick(r, "group", "Group", "ZHAI_matrix", "type") for r in rows],
        id=[_pick(r, "id", "No.", "Comment", "analysis") for r in rows],
    )


def load_groups(path):
    """IMPORT ZHAI/matrix assignment per analysis from a separate file (later).
    Expected: an id column + a group column (values like 'ZHAI' / 'matrix').
    Returns {id: group} or None if no file yet."""
    if not path or not os.path.exists(path):
        return None
    rows = _read_table(path)
    out = {}
    for r in rows:
        key = _pick(r, "id", "No.", "Comment", "analysis")
        grp = _pick(r, "group", "Group", "ZHAI_matrix", "type")
        if key is not None:
            out[str(key)] = grp
    return out


# ----------------------------------------------------------------------------
# I/O
# ----------------------------------------------------------------------------
def write_target(data, path):
    cols = ["Frac_Melt", "XCl_div_XOH", "XF_div_XOH", "XF_div_XCl"]
    n = len(data["Frac_Melt"])
    with open(path, "w", newline="", encoding="utf-8") as fh:
        w = csv.writer(fh)
        w.writerow(cols)
        for i in range(n):
            w.writerow(["%.15g" % float(data[c][i]) for c in cols])


def load_target(path):
    with open(path, newline="", encoding="utf-8-sig") as fh:
        rows = list(csv.DictReader(fh))
    F_obs = np.array([float(r["Frac_Melt"]) for r in rows])
    cloh = np.array([float(r["XCl_div_XOH"]) for r in rows])
    foh = np.array([float(r["XF_div_XOH"]) for r in rows])
    return F_obs, cloh, foh


def write_multioutput(results, path):
    with open(path, "w", newline="", encoding="utf-8") as fh:
        w = csv.writer(fh)
        w.writerow(MO_COLS)
        for rmse, x in results:
            w.writerow(["%.12g" % x[i] for i in range(len(PARAM_ORDER))] + ["%.12g" % rmse])


# ----------------------------------------------------------------------------
def _build_FT(num_frac):
    return am.build_F_and_T(TF_CURVE_FILE, CONST_T_C, num_frac)


def _self_check(F, T_K):
    """Vectorised fixed-point forward must match apatite_model's least_squares."""
    cloh_v, foh_v, conv = forward_ratios(TARGET_PARAMS, F, T_K, with_conv=True)
    if not conv:
        raise SystemExit("[self-check] fixed-point solve did NOT converge at TARGET_PARAMS")
    data, _ = am.forward_model(F, T_K, TARGET_PARAMS, K_SPECIATION)
    r1 = np.max(np.abs(cloh_v - data["XCl_div_XOH"]) / np.max(np.abs(data["XCl_div_XOH"])))
    r2 = np.max(np.abs(foh_v - data["XF_div_XOH"]) / np.max(np.abs(data["XF_div_XOH"])))
    worst = max(r1, r2)
    print(f"[self-check] fixed-point vs least_squares forward: max|rel| = {worst:.2e} "
          f"-> {'OK' if worst < 1e-9 else 'CHECK'}")
    return data


def cmd_create_target(args):
    F, T_K = _build_FT(NUM_FRAC_TARGET)
    print(f"[create-target] T: {T_K[0]-273.15:.0f} -> {T_K[-1]-273.15:.0f} degC; "
          f"{len(F)} steps")
    data = _self_check(F, T_K)
    out = args.target_out or os.path.join(HERE, "ot_target_curve.csv")
    write_target(data, out)
    print(f"[written] {out}")

    obs = load_observed_ratios(OBSERVED_RATIOS_FILE)
    if obs is None:
        print("[observed] OBSERVED_RATIOS_FILE not set -> skipped "
              "(set it to overplot/score the fit).")
    else:
        m = np.isfinite(obs["Cl_OH"]) & np.isfinite(obs["F_OH"])
        print(f"[observed] loaded {m.sum()} ratios from "
              f"{os.path.basename(OBSERVED_RATIOS_FILE)} "
              f"(grouping import {'ON' if GROUPS_FILE else 'pending'}).")


def cmd_multistart(args):
    tgt_path = args.target_out or os.path.join(HERE, "ot_target_curve.csv")
    if not os.path.exists(tgt_path):
        raise SystemExit(f"target curve not found ({tgt_path}); run create-target first.")
    F_obs, tgt_cloh, tgt_foh = load_target(tgt_path)
    F_fit, T_K_fit = _build_FT(NUM_FRAC_FIT)
    ns = args.num_start or NUM_START
    njobs = args.jobs if args.jobs is not None else (os.cpu_count() or 1)
    njobs = max(1, min(njobs, ns))
    print(f"[multistart] {ns} starts, {len(PARAM_ORDER)} params, fitting target "
          f"({len(F_obs)} pts) with {len(F_fit)}-pt forwards; {njobs} process(es) ...")
    results = run_multistart(F_obs, tgt_cloh, tgt_foh, F_fit, T_K_fit, ns, n_jobs=njobs)
    out = args.multi_out or os.path.join(HERE, "ot_multistart_bestfit.csv")
    write_multioutput(results, out)
    best = results[0]
    print(f"[best] RMSE={best[0]:.4g}")
    for k, v in zip(PARAM_ORDER, best[1]):
        print(f"       {k:10s} = {v:.4g}")
    print(f"[written] {out}  ({len(results)} fits, ascending RMSE)")


def main():
    ap = argparse.ArgumentParser(description="OT apatite Cl-OH-F inversion (Lormand workflow)")
    ap.add_argument("mode", choices=["create-target", "multistart", "both"])
    ap.add_argument("--num-start", type=int, default=None, help="override NUM_START")
    ap.add_argument("--jobs", type=int, default=None,
                    help="parallel processes for multistart (default: all cores; 1 = serial)")
    ap.add_argument("--target-out", default=None, help="target curve csv path")
    ap.add_argument("--multi-out", default=None, help="multistart output csv path")
    args = ap.parse_args()
    if args.mode in ("create-target", "both"):
        cmd_create_target(args)
    if args.mode in ("multistart", "both"):
        cmd_multistart(args)


if __name__ == "__main__":
    main()
