#!/usr/bin/env python3
# -*- coding: utf-8 -*-
r"""
Ok Tedi P1 -- apatite S -> melt S concentration tool (apatite_smelt).

WHAT THIS IS / IS NOT
---------------------
- Converts apatite S content (EPMA SO3) to MELT S *concentration* (ppm),
  an intensive quantity, via the Tassara et al. (2026) combination of the
  Konecke et al. (2019) fO2 dependence and Parat & Holtz (2004) T dependence
  (combined per Meng et al., 2021):

      Smelt (ppm) = S_Ap (ppm) * 10^( -0.263*dFMQ + 0.357 + 0.002*(T_C - 1000) )

- It is NOT a flux / tonnage / endowment calculator. Any extensive S-mass
  quantity is out of scope and is deliberately absent from this module.
- No sulfur valence (S6+/S2-) statement is made anywhere: P1 has no measured
  apatite S speciation. Code, comments and figure captions must stay silent
  on valence.

DELTA-FMQ FIREWALL
------------------
dFMQ (Loucks et al., 2020 zircon oxybarometer) is an INPUT of the conversion
formula. The computed Smelt must never be quoted as evidence of magma
oxidation state; the fO2 argument chain remains Loucks dFMQ + Ce4+/Ce3+ +
Eu/Eu* and is independent of this module. The guard sentence is stamped into
every output CSV header and into the Fig C caption.

TWO BRANCHES
------------
ZHAI (zircon-hosted apatite inclusions): each inclusion is paired one-to-one
  with its host zircon and uses that zircon's (dFMQ, T) directly. Inclusions
  lacking host data are dropped WITH a logged list (never silently).
  All points are computed regardless of host age; host_magmatic
  (host age <= 3 Ma, the Pliocene magmatic population) is carried as a flag,
  and ONLY magmatic-host points enter unit statistics / the Fig B main
  display (user ruling 2026-08-04: "compute but quarantine").

Matrix (phenocryst) apatite: per-sample "conversion-factor distribution"
  method (AE ruling 2026-08-04). For every screened zircon i of sample s the
  factor k_i = 10^(-0.263*dFMQ_i + 0.357 + 0.002*(T_i - 1000)) is computed
  from the PAIRED (dFMQ_i, T_i) of that grain, and the empirical distribution
  {k_i} gives k_med / k_P16 / k_P84 (numpy.percentile, linear interpolation,
  no parametric assumption). Then per apatite point:
      Smelt_med = S_Ap * k_med;  envelope = [S_Ap*k_P16, S_Ap*k_P84].

COVARIANCE RULE (hard constraint)
---------------------------------
Ti-in-zircon T and Loucks dFMQ share the Ti temperature term, so their errors
are negatively correlated within a sample. Computing k_i per PAIRED
(dFMQ_i, T_i) preserves that covariance. It is FORBIDDEN to take the medians
(or endpoints) of dFMQ and T separately and then combine them -- that breaks
the covariance structure and mis-states the envelope.

SENSITIVITY (for Methods)
-------------------------
d(log10 Smelt)/d(dFMQ) = -0.263  -> 1 log-unit dFMQ shifts Smelt by x/1.83.
d(log10 Smelt)/dT      = +0.002/degC -> 50 degC corresponds to ~26 %.
Tassara et al. (2026) state the absolute value is good to within ~2x; the
geologically meaningful quantity is the RELATIVE variation. Between-group
comparisons are robust where the between-group dFMQ contrast exceeds the
within-group spread. T is Ti-in-zircon; the calibration used is declared in
the manuscript Methods (Tassara's own T uses Crisp et al., 2023).

SCREENING AND EXCLUSIONS (user rulings 2026-08-04)
--------------------------------------------------
- Zircon k-pools use in_screened_set == "yes" (the user's manual screening)
  OR skarn-hosted zircons (samples 24OT02-43 / 24OT02-48; all 30 grains are
  1.1-2.2 Ma magmatic zircons, ruled usable). An age <= 3 Ma safeguard is
  applied to the whole pool (currently drops nothing).
- Edinburgh (24OT01-11) is excluded from P1: prep removes its rows with a
  loud log; the compute stage additionally hard-aborts if any Edinburgh
  sample id survives into its inputs.
- New York MD is a coeval, extra-mine-area intrusion. Its label is fixed to
  "coeval intrusion"; the words "unmineralized" / "barren" must never appear
  in any output, legend, log or comment.
- Alteration-affected samples (24OT02-43, 24OT02-01, 24OT02-10) are computed
  normally and flagged alteration_affected=True (no exclusion).
- SO3 <= 0 / missing = censored (below EPMA detection): no Smelt is computed
  and the point is only counted in the QC report. Substituting 0 or d.l./2
  is forbidden.

Outputs (apatite_smelt_out\): OT_smelt_points.csv, OT_smelt_sample_k.csv,
OT_smelt_unit_stats.csv, OT_smelt_qc.txt, Fig A/B/C/D as PDF+PNG.
Normalized input caches live in apatite_smelt_out\input\ (--reprep rebuilds).
Default output root: apatite_smelt_out\ beside this file for a script run;
a FROZEN app (PyInstaller) instead uses %OT_OUT%\apatite_smelt_out or
~\<exe-name>-output\apatite_smelt_out -- NEVER the onefile _MEIPASS
extraction dir, which the bootloader deletes on exit (2026-08-08 data loss).

Env: OT_SMELT_ZIRCON, OT_SMELT_ZHAI, OT_SMELT_MATRIX, OT_SMELT_GOLDEN,
     OT_SMELT_OUT override the default data paths / output dir.

References: Tassara et al. (2026) EPSL 690, 120139; Konecke et al. (2019);
Parat & Holtz (2004); Meng et al. (2021); Loucks et al. (2020) J. Petrol.
"""

from __future__ import annotations

import argparse
import csv
import datetime
import os
import re
import sys
import tempfile
import textwrap

import numpy as np

# ---------------------------------------------------------------------------
# constants (locked by spec p1-apatite-smelt-module.md + rulings 2026-08-04)
# ---------------------------------------------------------------------------
COEF_DFMQ = -0.263          # Konecke et al. (2019) fO2 dependence
COEF_INTERCEPT = 0.357
COEF_T = 0.002              # Parat & Holtz (2004) T dependence, per degC
T_REF_C = 1000.0
SO3_TO_S_PPM = 4005.0       # SO3 wt% -> S ppm: 1e4 * 32.065/80.063 (canonical)
                            # NOTE the reference Excel calculator used 4004.0
                            # (rel diff 2.5e-4); golden compare allows 3e-4.
N_MIN_ZIRCON = 8            # sample-level fallback threshold (spec default)
PERCENTILES = (16.0, 84.0)  # envelope percentiles
MAGMATIC_AGE_MA = 3.0       # Pliocene-population safeguard / ZHAI host flag
PROPAGATE_SIGMA_ZHAI = False  # spec 2.1 option, dormant: inputs carry no
                              # sigma_dfmq / sigma_t columns today

EDINBURGH_SAMPLE = "24OT01-11"

STAGE1 = "Stage1_Sydney"
STAGE2 = "Stage2_Kalgoorlie_Ningi"
STAGE3 = "Stage3_Fubilan"
COEVAL = "Coeval_NewYork"
SKARN = "Skarn_provisional"
UNIT_ORDER = [STAGE1, STAGE2, STAGE3, COEVAL, SKARN]

# canonical five-unit sample map (p1-intrusive-grouping, verified against
# "OT 数据表-syc.xlsx" 2026-08-04). Bare 24OT01-4 covers the EPMA dialect
# "24OT01--4"; syc splits it into 04a/04b/04c which all belong to the
# Fubilan suite, so the collapse is unit-safe. Fubilan 1/2/quartz-core are
# merged into Stage3 and must not be split.
SAMPLE_UNIT_MAP = {
    "24OT02-38": STAGE1, "24OT02-39": STAGE1,
    "24OT02-25": STAGE2, "24OT02-26": STAGE2, "24OT02-31": STAGE2,
    "24OT02-33": STAGE2, "24OT02-34": STAGE2, "24OT02-35": STAGE2,
    "24OT02-35-r": STAGE2,
    "24OT01-1": STAGE3, "24OT01-2": STAGE3, "24OT01-4": STAGE3,
    "24OT01-4a": STAGE3, "24OT01-4b": STAGE3, "24OT01-4c": STAGE3,
    "24OT02-1": STAGE3, "24OT02-3": STAGE3, "24OT02-4": STAGE3,
    "24OT02-6": STAGE3, "24OT02-10": STAGE3, "24OT02-14": STAGE3,
    "24OT02-16": STAGE3, "24OT02-20": STAGE3,
    "24OT02-57": COEVAL, "24OT02-59": COEVAL,
    "24OT02-43": SKARN, "24OT02-48": SKARN,
}

# display labels. New York is a coeval intrusion -- never "barren" etc.
UNIT_DISPLAY = {
    STAGE1: "Stage 1 (Sydney MD)",
    STAGE2: "Stage 2 (Kalgoorlie + Ningi)",
    STAGE3: "Stage 3 (Fubilan suite)",
    COEVAL: "New York (coeval intrusion)",
    SKARN: "Skarn-hosted (provisional)",
}

# palette follows the canonical 26.7.30 dFMQ-vs-T figure (Blues sequential
# for the three host-suite stages + grey for the coeval intrusion); skarn is
# drawn hollow/grey to signal its provisional attribution.
UNIT_COLOR = {
    STAGE1: "#9ecae1",
    STAGE2: "#4292c6",
    STAGE3: "#08519c",
    COEVAL: "#595959",
    SKARN: "#999999",
}
UNIT_HOLLOW = {SKARN}

# raw zircon-CSV unit strings -> canonical unit, used only as a QC
# cross-check against SAMPLE_UNIT_MAP (never as the mapping itself).
RAW_UNIT_HINT = {
    "Sydney MD": STAGE1,
    "Kalgoorlie": STAGE2,
    "Ningi": STAGE2,
    "Fubilan(1)": STAGE3,
    "Fubilan(2)": STAGE3,
    "Fubilan(Quartz core)": STAGE3,
    "New York monzondiorite": COEVAL,
    "Skarn+darai limestone": SKARN,
}

# mirror of apatite_m1.ALTERATION_SAMPLES (single-source rule):
# apatite_m1 imports tkinter at module level, so a headless CLI keeps a local
# mirror and self-checks against the original at startup (main()).
ALTERATION_SAMPLES = {"24OT02-43", "24OT02-01", "24OT02-10"}

GUARD_SENTENCE = ("dFMQ is an input term of the conversion formula; "
                  "Smelt is NOT independent evidence of magma oxidation "
                  "state. No apatite S valence is asserted.")
CAVEAT = ("intensive melt-S concentration (ppm) only; flux/tonnage/endowment "
          "quantities are out of scope")

_SCRIPT_DIR = os.path.dirname(os.path.abspath(__file__))

# Every progress line goes through log(); the CLI prints, the GUI panel swaps
# in its own sink via set_log_sink() so the text lands in its log pane.
_SINK = [print]


def log(msg):
    _SINK[0](msg)


def set_log_sink(fn):
    """Redirect log() output. fn=None restores print. Returns the previous."""
    prev = _SINK[0]
    _SINK[0] = print if fn is None else fn
    return prev


def _data_path(env, fname):
    """env override -> a file of that name next to this script."""
    v = os.environ.get(env)
    if v:
        return v
    return os.path.join(_SCRIPT_DIR, fname)


ZIRCON_CSV = _data_path("OT_SMELT_ZIRCON", "OT-P1-zircon-grouped.csv")
ZHAI_XLSX = _data_path("OT_SMELT_ZHAI", "Zircon-apa-OT.xlsx")
MATRIX_XLSX = _data_path("OT_SMELT_MATRIX", "ApThermo_OT.xlsx")
GOLDEN_XLSX = _data_path("OT_SMELT_GOLDEN",
                         "OkTedi_apatite_Smelt_calculator.xlsx")


def in_temp_dir(path):
    """True if path is rooted in the OS temp dir (case-folded, sep-safe).

    Catches the PyInstaller onefile extraction dir %TEMP%\\_MEIxxxxxx --
    which the bootloader DELETES on normal exit -- and stale copies of it
    restored from old .otproj files (2026-08-08 data-loss bug: a whole
    Melt S run vanished with the _MEI dir). Shared with the GUI panel.
    """
    try:
        p = os.path.normcase(os.path.abspath(path))
        t = os.path.normcase(os.path.abspath(tempfile.gettempdir()))
    except (OSError, ValueError):
        return False
    return p == t or p.startswith(t.rstrip(os.sep) + os.sep)


def _default_out_dir():
    """OT_SMELT_OUT env > script run: apatite_smelt_out\\ beside this file
    > frozen app: a VISIBLE per-user dir -- never sys._MEIPASS.

    Frozen (PyInstaller onefile) __file__/_SCRIPT_DIR resolve inside the
    %TEMP%\\_MEIxxxxxx extraction dir, so the pre-2026-08-08 default sent
    every Melt S run there and the bootloader deleted it with the app.
    Frozen default is %OT_OUT%\\apatite_smelt_out when OT_OUT is set (the
    app-wide output base), else ~\\<exe-name>-output\\apatite_smelt_out.
    A STABLE (non-timestamped) dir is required in all modes: the
    normalized input caches in <out>\\input\\ are reused across runs.
    """
    env = os.environ.get("OT_SMELT_OUT")
    if env:
        return env
    if getattr(sys, "frozen", False):
        base = os.environ.get("OT_OUT")
        if not base or in_temp_dir(base):
            exe = os.path.splitext(os.path.basename(sys.executable))[0]
            base = os.path.join(os.path.expanduser("~"), exe + "-output")
        return os.path.join(base, "apatite_smelt_out")
    return os.path.join(_SCRIPT_DIR, "apatite_smelt_out")


OUT_DIR = _default_out_dir()

MATRIX_SHEET = "Apatite stoichiometry-OT"
MATRIX_TYPE_KEEP = "ore-related-OT"     # phenocryst rows; "-zir" rows are the
                                        # ZHAI duplicates, other TYPEs are
                                        # other projects -- all excluded.


# ===========================================================================
# ==== pure computation (no I/O; unit-testable) ====
# ===========================================================================
_NUM_TOKEN = re.compile(r"^0*(\d+)([A-Za-z-]*)$")


def normalize_id(raw):
    """Canonicalize a sample/point id across the known dialects.

    "24OT01--11" -> "24OT01-11"; "24OT02-01" -> "24OT02-1";
    "24OT01-04a" -> "24OT01-4a"; first token is left untouched.
    """
    s = str(raw).strip()
    while "--" in s:
        s = s.replace("--", "-")
    toks = s.split("-")
    out = [toks[0]]
    for t in toks[1:]:
        m = _NUM_TOKEN.match(t)
        out.append((m.group(1) + m.group(2)) if m else t)
    return "-".join(out)


def split_zircon_id(raw):
    """Zircon-table ID dialects: '24OT01-1_10' / '24OT02-14 - 2'."""
    s = str(raw).strip()
    if "_" in s:
        head, grain = s.split("_", 1)
    elif " - " in s:
        head, grain = s.split(" - ", 1)
    else:
        head, grain = s, ""
    return normalize_id(head), grain.strip()


def split_zhai_id(raw):
    """'24OT02-59-02-zir' -> ('24OT02-59', '2')."""
    s = str(raw).strip()
    if s.lower().endswith("-zir"):
        s = s[:-4]
    s = normalize_id(s)
    head, _, grain = s.rpartition("-")
    return head, grain


def split_matrix_id(raw):
    """'24OT01--4-1' -> ('24OT01-4', '1'). Normalize FIRST (double hyphen)."""
    s = normalize_id(raw)
    head, _, grain = s.rpartition("-")
    return head, grain


def xf_xcl_ratio(xf, xcl):
    """Apatite X-site mole ratio XF/XCl from the sheets' own mole-fraction
    columns (computed here, not read from a derived column, mirroring the
    S_Ap-from-SO3 rule). None if either term is missing or XCl <= 0."""
    xf, xcl = to_float(xf), to_float(xcl)
    if xf is None or xcl is None or xcl <= 0:
        return None
    return xf / xcl


def to_float(x):
    """Tolerant float: None/''/text -> None; strips a trailing '?'."""
    if x is None:
        return None
    if isinstance(x, (int, float)):
        return float(x)
    s = str(x).strip().rstrip("?").strip()
    if not s:
        return None
    try:
        return float(s)
    except ValueError:
        return None


def k_factor(dfmq, t_c):
    """THE conversion factor; the formula lives only here."""
    return 10.0 ** (COEF_DFMQ * dfmq + COEF_INTERCEPT
                    + COEF_T * (t_c - T_REF_C))


def sap_ppm(so3_wt):
    """SO3 wt% -> S ppm; <=0 / missing -> None (censored; never 0 or d.l./2)."""
    v = to_float(so3_wt)
    if v is None or v <= 0.0:
        return None
    return v * SO3_TO_S_PPM


def k_pool_stats(pairs):
    """Empirical k distribution from paired (dfmq, T_C) tuples.

    Each k_i is computed from its own PAIRED (dfmq_i, T_i) -- covariance
    rule: never combine medians/endpoints of dfmq and T taken separately.
    """
    ks = [k_factor(d, t) for d, t in pairs]
    p16, med, p84 = np.percentile(ks, [PERCENTILES[0], 50.0, PERCENTILES[1]])
    return {
        "n": len(ks),
        "k_med": float(med), "k_p16": float(p16), "k_p84": float(p84),
        "dfmq_med": float(np.median([d for d, _ in pairs])),   # report-only
        "t_med": float(np.median([t for _, t in pairs])),      # report-only
    }


def unit_of(sample_id):
    return SAMPLE_UNIT_MAP[sample_id]


# ===========================================================================
# ==== prep stage: raw sources -> normalized input CSVs ====
# ===========================================================================
def _in_dir(out_dir):
    return os.path.join(out_dir, "input")


def _write_input_csv(path, rows, fieldnames, provenance):
    os.makedirs(os.path.dirname(path), exist_ok=True)
    with open(path, "w", encoding="utf-8", newline="") as f:
        for line in provenance:
            f.write(f"# {line}\n")
        w = csv.DictWriter(f, fieldnames=fieldnames)
        w.writeheader()
        for r in rows:
            w.writerow({k: ("" if r.get(k) is None else r[k])
                        for k in fieldnames})
    log(f"[written] {path} ({len(rows)} rows)")


def _read_input_csv(path):
    with open(path, "r", encoding="utf-8", newline="") as f:
        return list(csv.DictReader(ln for ln in f if not ln.startswith("#")))


def _cache_source(path):
    """The `# source: ...` provenance line of a normalized input cache
    (None when the file or the line is missing)."""
    try:
        with open(path, "r", encoding="utf-8") as f:
            for ln in f:
                if not ln.startswith("#"):
                    return None
                body = ln[1:].strip()
                if body.startswith("source:"):
                    return body.split("source:", 1)[1].strip()
    except OSError:
        return None
    return None


def _same_path(a, b):
    if not a or not b:
        return False
    try:
        return (os.path.normcase(os.path.abspath(a))
                == os.path.normcase(os.path.abspath(b)))
    except (OSError, ValueError):
        return False


# canonical normalized-input schemas (the smelt_input_*.csv fieldnames);
# used by run_prep both to WRITE the caches and to RECOGNISE an
# already-normalized table offered as a data source (user 2026-08-08)
NORM_FIELDS = {
    "zircon": ["zircon_id", "sample_id", "unit", "unit_src", "age_ma",
               "t_c", "dfmq", "screened"],
    "zhai": ["inclusion_id", "sample_id", "unit", "host_zircon_id",
             "host_age_ma", "so3_wt", "xf_xcl", "dfmq_host", "t_host_c"],
    "matrix": ["point_id", "sample_id", "unit", "so3_wt", "xf_xcl"],
}


def _csv_fieldnames(path):
    """Header row of a CSV (first non-# line), or None when unreadable."""
    try:
        with open(path, "r", encoding="utf-8-sig", newline="") as f:
            for ln in f:
                if ln.startswith("#"):
                    continue
                return [c.strip() for c in next(csv.reader([ln]))]
    except (OSError, UnicodeDecodeError, StopIteration, csv.Error):
        return None
    return None


def _looks_normalized(src, tag):
    """True when src already carries the normalized smelt_input_<tag>
    schema (header ⊇ NORM_FIELDS[tag]) — e.g. an edited copy of a
    normalized cache offered as a data source. Such files bypass the raw
    prep and are ingested as-is (rows trusted, incl. their `unit` column;
    the compute-stage guards still apply)."""
    if not str(src).lower().endswith(".csv"):
        return False
    cols = _csv_fieldnames(src)
    return cols is not None and set(NORM_FIELDS[tag]) <= set(cols)


def _provenance(src, n_raw, n_kept, extra=()):
    return [f"Ok Tedi P1 apatite_smelt normalized input. CAVEAT: {CAVEAT}",
            f"source: {src}",
            f"prepared: {datetime.date.today().isoformat()}  "
            f"raw rows={n_raw} kept={n_kept}",
            *extra]


def _drop_edinburgh(rows, key, label):
    kept, dropped = [], []
    for r in rows:
        (dropped if r[key] == EDINBURGH_SAMPLE else kept).append(r)
    if dropped:
        ids = ", ".join(str(r.get("point_id") or r.get("inclusion_id")
                            or r.get("zircon_id")) for r in dropped)
        log(f"[edinburgh] {label}: excluded {len(dropped)} row(s) of sample "
              f"{EDINBURGH_SAMPLE} (excluded from P1, canonical grouping): "
              f"{ids}")
    return kept, len(dropped)


def _assert_mapped(rows, label):
    missing = sorted({r["sample_id"] for r in rows
                     if r["sample_id"] not in SAMPLE_UNIT_MAP})
    if missing:
        sys.exit(f"[FATAL] {label}: unmapped sample id(s) -- extend "
                 f"SAMPLE_UNIT_MAP deliberately, do not guess: "
                 f"{', '.join(missing)}")


def prep_zircon(src):
    with open(src, "r", encoding="utf-8-sig", newline="") as f:
        raw = list(csv.DictReader(f))
    if not raw or "ID" not in raw[0]:
        got = ", ".join(raw[0].keys()) if raw else "(empty file)"
        sys.exit(f"[FATAL] zircon source {src}: expected the RAW grouped "
                 f"CSV (needs an 'ID' column) or a normalized "
                 f"smelt_input_zircon table (columns "
                 f"{', '.join(NORM_FIELDS['zircon'])}); got columns: {got}")
    rows = []
    for r in raw:
        sample, grain = split_zircon_id(r["ID"])
        rows.append({
            "zircon_id": r["ID"].strip(), "sample_id": sample,
            "unit_src": r["unit"].strip(),
            "age_ma": to_float(r["age_Ma"]),
            "t_c": to_float(r["T_Ti_in_zircon_C"]),
            "dfmq": to_float(r["dFMQ"]),
            "screened": r["in_screened_set"].strip().lower(),
        })
    rows, n_ed = _drop_edinburgh(rows, "sample_id", "zircon")
    _assert_mapped(rows, "zircon")
    for r in rows:
        r["unit"] = unit_of(r["sample_id"])
        hint = RAW_UNIT_HINT.get(r["unit_src"])
        if hint is not None and hint != r["unit"]:
            log(f"[WARNING] zircon {r['zircon_id']}: raw unit "
                  f"'{r['unit_src']}' conflicts with SAMPLE_UNIT_MAP "
                  f"({r['unit']})")
    return rows, len(raw), n_ed


def _require_xlsx(src, tag):
    if not str(src).lower().endswith((".xlsx", ".xlsm")):
        sys.exit(f"[FATAL] {tag} source {src}: expected the RAW Excel "
                 f"workbook or a normalized smelt_input_{tag} CSV (columns "
                 f"{', '.join(NORM_FIELDS[tag])}).")


def prep_zhai(src):
    _require_xlsx(src, "zhai")
    from openpyxl import load_workbook
    ws = load_workbook(src, data_only=True, read_only=True)["Sheet2"]
    it = ws.iter_rows(values_only=True)
    hdr = [str(c).strip() if c is not None else "" for c in next(it)]
    col = {name: hdr.index(name) for name in
           ("SAMPLE", "SO3", "XF", "XCl", "zir-age", "Tzir", "FMQ")}
    raw = [r for r in it if r[col["SAMPLE"]] not in (None, "")]
    # blank-SAMPLE rows are MIN/MAX/AVE summary blocks -- already filtered.
    rows = []
    for r in raw:
        incl = str(r[col["SAMPLE"]]).strip()
        sample, grain = split_zhai_id(incl)
        rows.append({
            "inclusion_id": incl, "sample_id": sample,
            "host_zircon_id": f"{sample}-{grain}",
            "host_age_ma": to_float(r[col["zir-age"]]),
            "so3_wt": to_float(r[col["SO3"]]),
            "xf_xcl": xf_xcl_ratio(r[col["XF"]], r[col["XCl"]]),
            "dfmq_host": to_float(r[col["FMQ"]]),
            "t_host_c": to_float(r[col["Tzir"]]),
        })
    rows, n_ed = _drop_edinburgh(rows, "sample_id", "zhai")
    _assert_mapped(rows, "zhai")
    for r in rows:
        r["unit"] = unit_of(r["sample_id"])
    return rows, len(raw), n_ed


def prep_matrix(src):
    _require_xlsx(src, "matrix")
    from openpyxl import load_workbook
    ws = load_workbook(src, data_only=True, read_only=True)[MATRIX_SHEET]
    grid = list(ws.iter_rows(values_only=True))
    hrow = next(i for i, r in enumerate(grid)
                if r and str(r[0]).strip() == "TYPE")
    hdr = [str(c).strip() if c is not None else "" for c in grid[hrow]]
    i_type, i_sample = hdr.index("TYPE"), hdr.index("SAMPLE")
    i_so3 = hdr.index("SO3")            # .index -> FIRST match; the sheet has
    assert i_so3 == 15, f"matrix SO3 column moved (got {i_so3}, want 15)"
    # a second 'SO3' at index 37 (apfu block) that must not be picked up.
    i_xf, i_xcl = hdr.index("XF"), hdr.index("XCl")
    i_fclr = hdr.index("XF/XCL")        # precomputed ratio: cross-check only
    assert (i_xf, i_xcl, i_fclr) == (20, 21, 25), \
        f"matrix XF/XCl columns moved (got {(i_xf, i_xcl, i_fclr)}, want (20, 21, 25))"
    raw = [r for r in grid[hrow + 1:]
           if r and r[i_type] is not None and r[i_sample] is not None]
    rows = []
    ratio_dev = 0.0
    for r in raw:
        if str(r[i_type]).strip() != MATRIX_TYPE_KEEP:
            continue
        pid = str(r[i_sample]).strip()
        sample, grain = split_matrix_id(pid)
        ratio = xf_xcl_ratio(r[i_xf], r[i_xcl])
        pre = to_float(r[i_fclr])
        if ratio is not None and pre is not None and pre > 0:
            ratio_dev = max(ratio_dev, abs(ratio - pre) / pre)
        rows.append({"point_id": pid, "sample_id": sample,
                     "so3_wt": to_float(r[i_so3]), "xf_xcl": ratio})
    n_kept_type = len(rows)
    if ratio_dev > 1e-6:
        log(f"[WARNING] matrix XF/XCl: computed ratio deviates from the "
              f"sheet's XF/XCL column (max rel {ratio_dev:.2e})")
    else:
        log(f"[input] matrix XF/XCl cross-check vs sheet column: "
              f"max rel diff {ratio_dev:.2e}")
    rows, n_ed = _drop_edinburgh(rows, "sample_id", "matrix")
    _assert_mapped(rows, "matrix")
    for r in rows:
        r["unit"] = unit_of(r["sample_id"])
    return rows, n_kept_type, n_ed


def run_prep(out_dir, force, sources=None):
    """sources: optional {tag: path} override of the three raw inputs
    ("zircon"/"zhai"/"matrix"; GUI data-source dialog, 2026-08-08); any
    missing key falls back to the canonical module constant. A cached
    normalized input is reused only when its recorded `# source:` line
    matches the requested file — switching sources in EITHER direction
    rebuilds automatically instead of silently mixing datasets (this also
    retires the old footgun where a cache built from different raw files
    was reused as long as it merely existed)."""
    sources = sources or {}
    src_map = {"zircon": sources.get("zircon") or ZIRCON_CSV,
               "zhai": sources.get("zhai") or ZHAI_XLSX,
               "matrix": sources.get("matrix") or MATRIX_XLSX}
    ind = _in_dir(out_dir)
    paths = {t: os.path.join(ind, f"smelt_input_{t}.csv")
             for t in ("zircon", "zhai", "matrix")}
    if not force and all(os.path.exists(p) for p in paths.values()):
        stale = [t for t in ("zircon", "zhai", "matrix")
                 if not _same_path(_cache_source(paths[t]), src_map[t])]
        if not stale:
            log(f"[input] cached normalized inputs in {ind} (use --reprep to "
                  f"rebuild)")
            for t in ("zircon", "zhai", "matrix"):
                log(f"[input] {t} (cached) <- {src_map[t]}")
            return paths
        log(f"[input] cache/source mismatch ({', '.join(stale)}) -- "
            f"rebuilding the normalized inputs from the requested sources")
    for tag in ("zircon", "zhai", "matrix"):
        if not os.path.exists(src_map[tag]):
            sys.exit(f"[FATAL] input source missing: {src_map[tag]}\n  set "
                     f"OT_SMELT_{tag.upper()} to the correct file (or pick "
                     f"it in the GUI Data sources dialog).")
    preps = {"zircon": prep_zircon, "zhai": prep_zhai, "matrix": prep_matrix}
    extra = {"zircon": [], "zhai": [],
             "matrix": [f"TYPE filter: {MATRIX_TYPE_KEEP}"]}
    for tag in ("zircon", "zhai", "matrix"):
        src = src_map[tag]
        if _looks_normalized(src, tag):
            # already-normalized table offered as a source (user 2026-08-08):
            # ingest as-is, prep skipped; its `unit` column is trusted and
            # the compute-stage guards (Edinburgh, uniqueness) still apply
            rows = _read_input_csv(src)
            unknown = sorted({r["unit"] for r in rows
                              if r["unit"] not in UNIT_DISPLAY})
            if unknown:
                log(f"[input] {tag}: NOTE unit value(s) outside the "
                    f"canonical set (rows kept; they will not appear in the "
                    f"canonical unit ordering/statistics): "
                    f"{', '.join(unknown)}")
            log(f"[input] {tag} <- {src} (normalized-format table, prep "
                f"skipped)")
            _write_input_csv(paths[tag], rows, NORM_FIELDS[tag],
                             _provenance(src, len(rows), len(rows),
                                         ["normalized-format import: prep "
                                          "skipped, rows taken as-is"]))
            continue
        log(f"[input] {tag} <- {src}")
        rows, n_raw, n_ed = preps[tag](src)
        _write_input_csv(paths[tag], rows, NORM_FIELDS[tag],
                         _provenance(src, n_raw, len(rows),
                                     [f"edinburgh excluded: {n_ed}"]
                                     + extra[tag]))
    return paths


# ===========================================================================
# ==== compute stage (consumes normalized CSVs only) ====
# ===========================================================================
def guard_edinburgh(tables):
    """Hard abort if any Edinburgh sample id survives into compute inputs.

    Uses the parsed sample_id -- NEVER substring matching (nine samples
    carry a grain '-11' in their point ids).
    """
    for label, rows in tables.items():
        hit = [r for r in rows if r["sample_id"] == EDINBURGH_SAMPLE]
        if hit:
            sys.exit(f"[FATAL] Edinburgh sample {EDINBURGH_SAMPLE} detected "
                     f"in {label} input ({len(hit)} rows). Edinburgh is "
                     f"excluded from P1; fix the input preparation.")


def _assert_unique_points(rows, key, label):
    seen = {}
    for r in rows:
        k = normalize_id(r[key])
        if k in seen:
            sys.exit(f"[FATAL] {label}: normalized id collision '{k}' "
                     f"({seen[k]} vs {r[key]}) -- would double-count.")
        seen[k] = r[key]


def screen_zircons(zircon):
    """k-pool screening: manual screened set, plus skarn magmatic zircons
    (ruling 2026-08-04), with an age <= MAGMATIC_AGE_MA safeguard."""
    kept, dropped_age = [], []
    for r in zircon:
        take = (r["screened"] == "yes") or (r["unit"] == SKARN)
        if not take:
            continue
        age = to_float(r["age_ma"])
        if age is not None and age > MAGMATIC_AGE_MA:
            dropped_age.append(r["zircon_id"])
            continue
        kept.append(r)
    if dropped_age:
        log(f"[WARNING] age safeguard (> {MAGMATIC_AGE_MA} Ma) removed "
              f"{len(dropped_age)} screened zircon(s): "
              f"{', '.join(dropped_age)}")
    n_yes = sum(1 for r in kept if r["screened"] == "yes")
    log(f"[screen] k-pool zircons: {len(kept)} "
          f"(manual screened {n_yes} + skarn ruling {len(kept) - n_yes})")
    return kept


def build_k_pools(screened):
    by_sample, by_unit = {}, {}
    for r in screened:
        d, t = to_float(r["dfmq"]), to_float(r["t_c"])
        if d is None or t is None:
            log(f"[WARNING] screened zircon {r['zircon_id']} lacks "
                  f"dfmq/t_c -- skipped from pools")
            continue
        by_sample.setdefault(r["sample_id"], []).append((d, t))
        by_unit.setdefault(r["unit"], []).append((d, t))
    return by_sample, by_unit


def compute_zhai_points(zhai):
    points, unpaired = [], []
    for r in zhai:
        dfmq, t = to_float(r["dfmq_host"]), to_float(r["t_host_c"])
        if dfmq is None or t is None:
            unpaired.append(r["inclusion_id"])
            continue                       # dropped, logged -- never silent
        age = to_float(r["host_age_ma"])
        sap = sap_ppm(r["so3_wt"])
        censored = sap is None
        k = k_factor(dfmq, t)
        points.append({
            "type": "ZHAI", "point_id": r["inclusion_id"],
            "sample_id": r["sample_id"], "unit": r["unit"],
            "s_ap_ppm": sap, "xf_xcl": to_float(r["xf_xcl"]),
            "smelt_med": None if censored else sap * k,
            "smelt_p16": None, "smelt_p84": None,   # per spec: no envelope
            "k_med": None, "k_p16": None, "k_p84": None,
            "fallback": "", "censored": censored,
            "alteration_affected": _is_alteration(r["sample_id"]),
            "host_zircon_id": r["host_zircon_id"],
            "host_age_ma": age,
            "host_magmatic": (age is not None and age <= MAGMATIC_AGE_MA),
            "host_dfmq": dfmq, "host_t_c": t,
        })
    if unpaired:
        log(f"[unpaired] {len(unpaired)} ZHAI point(s) lack host zircon "
              f"data and were dropped: {', '.join(unpaired)}")
    return points, unpaired


def compute_matrix_points(matrix, by_sample, by_unit):
    samples = sorted({(r["sample_id"], r["unit"]) for r in matrix})
    sample_k, starved = {}, []
    for sid, unit in samples:
        own = by_sample.get(sid, [])            # n = 0 (no key) is a valid
        if len(own) >= N_MIN_ZIRCON:            # fallback path, not an error
            st, fb = k_pool_stats(own), False
        else:
            pool = by_unit.get(unit, [])
            if len(pool) < N_MIN_ZIRCON:
                starved.append(f"{sid} (unit {unit}: n={len(pool)})")
                continue
            st, fb = k_pool_stats(pool), True
            log(f"[fallback] {sid}: n_zircon={len(own)} < {N_MIN_ZIRCON} "
                  f"-> unit pool {unit} (n={len(pool)})")
        sample_k[sid] = {"sample_id": sid, "unit": unit,
                         "n_zircon": len(own), "pool_n": st["n"],
                         "fallback": fb, **st}
    if starved:
        sys.exit("[FATAL] unit-level k-pool still below n_min for: "
                 + "; ".join(starved))
    points = []
    for r in matrix:
        st = sample_k[r["sample_id"]]
        sap = sap_ppm(r["so3_wt"])
        censored = sap is None
        points.append({
            "type": "matrix", "point_id": r["point_id"],
            "sample_id": r["sample_id"], "unit": r["unit"],
            "s_ap_ppm": sap, "xf_xcl": to_float(r["xf_xcl"]),
            "smelt_med": None if censored else sap * st["k_med"],
            "smelt_p16": None if censored else sap * st["k_p16"],
            "smelt_p84": None if censored else sap * st["k_p84"],
            "k_med": st["k_med"], "k_p16": st["k_p16"], "k_p84": st["k_p84"],
            "fallback": st["fallback"], "censored": censored,
            "alteration_affected": _is_alteration(r["sample_id"]),
            "host_zircon_id": None, "host_age_ma": None,
            "host_magmatic": None, "host_dfmq": None, "host_t_c": None,
        })
    return points, sample_k


def _is_alteration(sample_id):
    return normalize_id(sample_id) in {normalize_id(s)
                                       for s in ALTERATION_SAMPLES}


def _stats_pool(points, typ, unit):
    """Points entering unit statistics: non-censored; ZHAI additionally
    restricted to magmatic hosts (ruling: compute-but-quarantine)."""
    return [p["smelt_med"] for p in points
            if p["type"] == typ and p["unit"] == unit
            and not p["censored"]
            and (typ != "ZHAI" or p["host_magmatic"])]


def unit_stats(points):
    rows = []
    for unit in UNIT_ORDER:
        for typ in ("ZHAI", "matrix"):
            vals = _stats_pool(points, typ, unit)
            n_cen = sum(1 for p in points if p["type"] == typ
                        and p["unit"] == unit and p["censored"])
            n_inh = sum(1 for p in points if p["type"] == typ
                        and p["unit"] == unit and typ == "ZHAI"
                        and not p["censored"] and not p["host_magmatic"])
            if not vals:
                log(f"[WARNING] unit stats: no usable {typ} points for "
                      f"{unit} -- row emitted with n=0")
                rows.append({"unit": unit, "type": typ, "n": 0,
                             "smelt_median": None, "smelt_p16": None,
                             "smelt_p84": None, "n_censored": n_cen,
                             "n_inherited_host_excluded": n_inh})
                continue
            p16, med, p84 = np.percentile(
                vals, [PERCENTILES[0], 50.0, PERCENTILES[1]])
            rows.append({"unit": unit, "type": typ, "n": len(vals),
                         "smelt_median": float(med),
                         "smelt_p16": float(p16), "smelt_p84": float(p84),
                         "n_censored": n_cen,
                         "n_inherited_host_excluded": n_inh})
    return rows


# ===========================================================================
# ==== outputs / QC ====
# ===========================================================================
def _fmt(v):
    if v is None or v == "":
        return ""
    if isinstance(v, bool):
        return str(v)
    if isinstance(v, float):
        return f"{v:.6g}"
    return v


def _csv_header_lines(title):
    return [
        f"Ok Tedi P1 apatite_smelt -- {title}. CAVEAT: {CAVEAT}",
        f"GUARD: {GUARD_SENTENCE}",
        f"model: Smelt = S_Ap * 10^({COEF_DFMQ}*dFMQ + {COEF_INTERCEPT} + "
        f"{COEF_T}*(T-{T_REF_C:g}))  [Tassara et al. 2026]",
        f"SO3->S ppm factor {SO3_TO_S_PPM:g} (canonical); n_min="
        f"{N_MIN_ZIRCON}; envelope percentiles P{PERCENTILES[0]:g}/P"
        f"{PERCENTILES[1]:g}; magmatic host age <= {MAGMATIC_AGE_MA:g} Ma",
        f"generated: {datetime.date.today().isoformat()}",
    ]


def _write_csv(path, rows, fieldnames, title):
    with open(path, "w", encoding="utf-8", newline="") as f:
        for line in _csv_header_lines(title):
            f.write(f"# {line}\n")
        w = csv.DictWriter(f, fieldnames=fieldnames)
        w.writeheader()
        for r in rows:
            w.writerow({k: _fmt(r.get(k)) for k in fieldnames})
    log(f"[written] {path} ({len(rows)} rows)")


POINT_FIELDS = ["type", "point_id", "sample_id", "unit", "s_ap_ppm",
                "xf_xcl",
                "smelt_med", "smelt_p16", "smelt_p84", "k_med", "k_p16",
                "k_p84", "fallback", "censored", "alteration_affected",
                "host_zircon_id", "host_age_ma", "host_magmatic",
                "host_dfmq", "host_t_c"]


def write_outputs(out_dir, points, sample_k, ustats):
    _write_csv(os.path.join(out_dir, "OT_smelt_points.csv"), points,
               POINT_FIELDS, "per-point melt S")
    sk = [sample_k[s] for s in sorted(sample_k)]
    _write_csv(os.path.join(out_dir, "OT_smelt_sample_k.csv"), sk,
               ["sample_id", "unit", "n_zircon", "pool_n", "k_med", "k_p16",
                "k_p84", "dfmq_med", "t_med", "fallback"],
               "per-sample conversion factors")
    _write_csv(os.path.join(out_dir, "OT_smelt_unit_stats.csv"), ustats,
               ["unit", "type", "n", "smelt_median", "smelt_p16",
                "smelt_p84", "n_censored", "n_inherited_host_excluded"],
               "unit-level summary (ZHAI magmatic-host only)")


def write_qc(out_dir, meta, points, sample_k, unpaired):
    lines = [f"Ok Tedi P1 apatite_smelt QC report -- "
             f"{datetime.date.today().isoformat()}",
             f"CAVEAT: {CAVEAT}", f"GUARD: {GUARD_SENTENCE}", ""]
    lines.append("1) qualified zircons per sample / fallback events:")
    for sid in sorted(sample_k):
        st = sample_k[sid]
        tag = (f"FALLBACK -> unit pool {st['unit']} (n={st['pool_n']})"
               if st["fallback"] else "sample-level")
        lines.append(f"   {sid:14s} n_zircon={st['n_zircon']:3d}  {tag}")
    lines.append("")
    lines.append("2) censored analyses (SO3 <= d.l.; no Smelt computed):")
    for typ in ("ZHAI", "matrix"):
        cen = [p for p in points if p["type"] == typ and p["censored"]]
        by = {}
        for p in cen:
            by.setdefault(p["sample_id"], []).append(p["point_id"])
        lines.append(f"   {typ}: {len(cen)} point(s)")
        for sid in sorted(by):
            lines.append(f"      {sid}: {', '.join(by[sid])}")
    lines.append("")
    lines.append(f"3) unpaired ZHAI (no host zircon data; dropped): "
                 f"{len(unpaired)}")
    for u in unpaired:
        lines.append(f"      {u}")
    lines.append("")
    lines.append("4) unmapped samples: none (hard-error guard passed)")
    lines.append(f"5) Edinburgh: excluded at prep "
                 f"(zircon {meta['ed_zircon']}, zhai {meta['ed_zhai']}, "
                 f"matrix {meta['ed_matrix']} rows); compute guard passed")
    lines.append("")
    n_alt = sum(1 for p in points if p["alteration_affected"])
    lines.append(f"6) alteration_affected flagged points: {n_alt} "
                 f"(samples {sorted(ALTERATION_SAMPLES)}; flagged, not "
                 f"excluded)")
    zh = [p for p in points if p["type"] == "ZHAI"]
    n_mag = sum(1 for p in zh if p["host_magmatic"])
    lines.append(f"7) ZHAI host split: magmatic {n_mag} / inherited "
                 f"{len(zh) - n_mag} (inherited computed but quarantined "
                 f"from unit stats, Fig B main display, and Fig D)")
    path = os.path.join(out_dir, "OT_smelt_qc.txt")
    with open(path, "w", encoding="utf-8") as f:
        f.write("\n".join(lines) + "\n")
    log(f"[written] {path}")


# ===========================================================================
# ==== figures ====
# The draw_fig_* functions paint onto a caller-supplied Axes and touch no
# backend and no disk, so the GUI panel (apatite_smelt_panel.py) can render
# the same figures into an embedded canvas. The fig_* wrappers below are the
# CLI path: they own the Agg backend and the PDF+PNG save.
# ===========================================================================
FIG_KEYS = ("A", "B", "C", "D")
FIG_TITLES = {
    "A": "Fig A  apatite S vs dFMQ -- behavior check (inputs only)",
    "B": ("Fig B  melt S by intrusive unit "
          "(circles: matrix apatite; triangles: ZHAI)"),
    "C": "Fig C  melt S vs dFMQ (descriptive panel)",
    "D": ("Fig D  melt S vs apatite XF/XCl (boxes: per-unit Smelt, whis "
          "5-95%; circles matrix, triangles ZHAI)"),
}
FIG_STEMS = {"A": "OT_smelt_figA_sap_vs_dfmq",
             "B": "OT_smelt_figB_units",
             "C": "OT_smelt_figC_smelt_vs_dfmq",
             "D": "OT_smelt_figD_smelt_vs_xfxcl"}


def draw_figure(key, ax, points, sample_k, footnote=True):
    """Dispatch by FIG_KEYS; the single entry point the GUI panel uses."""
    {"A": draw_fig_a, "B": draw_fig_b, "C": draw_fig_c,
     "D": draw_fig_d}[key](ax, points, sample_k, footnote=footnote)


def _save_fig(fig, out_dir, stem):
    pdf = os.path.join(out_dir, stem + ".pdf")
    png = os.path.join(out_dir, stem + ".png")
    fig.savefig(pdf, bbox_inches="tight")
    fig.savefig(png, dpi=200, bbox_inches="tight")
    log(f"[written] {pdf}")
    log(f"[written] {png}")


def _footnote(ax, note, on=True):
    if not on:
        return
    ax.text(0.5, -0.17, "\n".join(textwrap.wrap(note, 95)),
            transform=ax.transAxes, ha="center", va="top",
            fontsize=6.5, color="#555555")


def _marker_kw(unit, filled_color=None):
    col = filled_color or UNIT_COLOR[unit]
    if unit in UNIT_HOLLOW:
        return dict(facecolors="none", edgecolors=col, linewidths=0.9)
    return dict(facecolors=col, edgecolors="#333333", linewidths=0.4)


def draw_fig_a(ax, points, sample_k, footnote=True):
    """Fig A: measured S_Ap vs dFMQ (behavior check reproduction).

    ZHAI use host-zircon dFMQ; matrix apatite use their sample's dfmq_med.
    This is an independent behavior check previously confirmed; it is NOT
    part of the Smelt conversion chain."""
    for unit in UNIT_ORDER:
        zh = [p for p in points if p["type"] == "ZHAI" and p["unit"] == unit
              and not p["censored"]]
        mag = [p for p in zh if p["host_magmatic"]]
        inh = [p for p in zh if not p["host_magmatic"]]
        if mag:
            ax.scatter([p["host_dfmq"] for p in mag],
                       [p["s_ap_ppm"] for p in mag], marker="^", s=34,
                       zorder=3, label=f"{UNIT_DISPLAY[unit]} ZHAI",
                       **_marker_kw(unit))
        if inh:
            ax.scatter([p["host_dfmq"] for p in inh],
                       [p["s_ap_ppm"] for p in inh], marker="^", s=22,
                       facecolors="none", edgecolors="0.6", linewidths=0.7,
                       zorder=2)
        mx = [p for p in points if p["type"] == "matrix"
              and p["unit"] == unit and not p["censored"]]
        if mx:
            ax.scatter([sample_k[p["sample_id"]]["dfmq_med"] for p in mx],
                       [p["s_ap_ppm"] for p in mx], marker="o", s=22,
                       alpha=0.85, zorder=3,
                       label=f"{UNIT_DISPLAY[unit]} matrix",
                       **_marker_kw(unit))
    ax.set_yscale("log")
    ax.set_xlabel("zircon dFMQ (ZHAI: host grain; matrix: sample median)",
                  fontsize=9)
    ax.set_ylabel("S in apatite (ppm)", fontsize=9)
    ax.tick_params(labelsize=7)
    ax.set_title(FIG_TITLES["A"], fontsize=9)
    ax.legend(fontsize=6.2, ncol=2, frameon=False)
    _footnote(ax, "Reproduction of the previously confirmed cross-check of "
                  "measured apatite S against zircon dFMQ. Grey open "
                  "triangles: ZHAI with inherited (pre-Pliocene) host "
                  "zircons. Censored (SO3 <= d.l.) points omitted. "
                  + GUARD_SENTENCE, footnote)


def draw_fig_b(ax, points, sample_k=None, footnote=True):
    """Fig B (main): Smelt by canonical unit, box + jittered scatter;
    ZHAI vs matrix by marker; matrix points carry the k_P16-P84 envelope;
    inherited-host ZHAI shown de-emphasized and excluded from boxes."""
    rng = np.random.default_rng(42)      # fixed: jitter must be reproducible
    for i, unit in enumerate(UNIT_ORDER):
        for typ, dx, marker in (("matrix", -0.18, "o"), ("ZHAI", 0.18, "^")):
            vals = _stats_pool(points, typ, unit)
            if len(vals) >= 3:
                bp = ax.boxplot([vals], positions=[i + dx], widths=0.26,
                                whis=(5, 95), showfliers=False,
                                patch_artist=True, zorder=1)
                bp["boxes"][0].set(facecolor=UNIT_COLOR[unit], alpha=0.18,
                                   edgecolor="#666666", linewidth=0.8)
                for part in ("medians", "whiskers", "caps"):
                    for art in bp[part]:
                        art.set(color="#666666", linewidth=0.8)
                bp["medians"][0].set(color="#222222", linewidth=1.2)
            pts = [p for p in points if p["type"] == typ
                   and p["unit"] == unit and not p["censored"]
                   and (typ != "ZHAI" or p["host_magmatic"])]
            if pts:
                xs = i + dx + rng.uniform(-0.06, 0.06, len(pts))
                ys = [p["smelt_med"] for p in pts]
                if typ == "matrix":
                    yerr = [[p["smelt_med"] - p["smelt_p16"] for p in pts],
                            [p["smelt_p84"] - p["smelt_med"] for p in pts]]
                    ax.errorbar(xs, ys, yerr=yerr, fmt="none",
                                ecolor=UNIT_COLOR[unit], elinewidth=0.5,
                                alpha=0.45, zorder=2)
                ax.scatter(xs, ys, marker=marker, s=26 if typ == "ZHAI"
                           else 18, zorder=3, **_marker_kw(unit))
            if typ == "ZHAI":
                inh = [p for p in points if p["type"] == "ZHAI"
                       and p["unit"] == unit and not p["censored"]
                       and not p["host_magmatic"]]
                if inh:
                    xs = i + dx + rng.uniform(-0.06, 0.06, len(inh))
                    ax.scatter(xs, [p["smelt_med"] for p in inh], marker="^",
                               s=18, facecolors="none", edgecolors="0.65",
                               linewidths=0.7, zorder=2)
    ax.set_yscale("log")
    ax.set_xticks(range(len(UNIT_ORDER)))
    ax.set_xticklabels([UNIT_DISPLAY[u].replace(" (", "\n(")
                        for u in UNIT_ORDER], fontsize=7)
    ax.set_ylabel("melt S (ppm)", fontsize=9)
    ax.tick_params(axis="y", labelsize=7)
    ax.set_title(FIG_TITLES["B"], fontsize=9)
    _footnote(ax, "Matrix-apatite error bars span the sample k P16-P84 "
                  "envelope (paired-zircon conversion-factor distribution). "
                  "Boxes: median, P25-P75, 5-95 whiskers of the shown "
                  "points. Grey open triangles: ZHAI with inherited "
                  "(pre-Pliocene) hosts -- computed but excluded from boxes "
                  "and unit statistics. Absolute values good to ~2x "
                  "(Tassara et al., 2026); relative between-unit contrasts "
                  "are the meaningful quantity. " + GUARD_SENTENCE, footnote)


def draw_fig_c(ax, points, sample_k, footnote=True):
    """Fig C (optional, descriptive): Smelt vs dFMQ. The caption MUST carry
    the guard sentence -- dFMQ is an input of the conversion."""
    for unit in UNIT_ORDER:
        mag = [p for p in points if p["type"] == "ZHAI" and p["unit"] == unit
               and not p["censored"] and p["host_magmatic"]]
        if mag:
            ax.scatter([p["host_dfmq"] for p in mag],
                       [p["smelt_med"] for p in mag], marker="^", s=34,
                       zorder=3, label=f"{UNIT_DISPLAY[unit]} ZHAI",
                       **_marker_kw(unit))
        mx = [p for p in points if p["type"] == "matrix"
              and p["unit"] == unit and not p["censored"]]
        if mx:
            ax.scatter([sample_k[p["sample_id"]]["dfmq_med"] for p in mx],
                       [p["smelt_med"] for p in mx], marker="o", s=22,
                       alpha=0.85, zorder=3,
                       label=f"{UNIT_DISPLAY[unit]} matrix",
                       **_marker_kw(unit))
    ax.set_yscale("log")
    ax.set_xlabel("dFMQ (input term of the conversion)", fontsize=9)
    ax.set_ylabel("melt S (ppm)", fontsize=9)
    ax.tick_params(labelsize=7)
    ax.set_title(FIG_TITLES["C"], fontsize=9)
    ax.legend(fontsize=6.2, ncol=2, frameon=False)
    _footnote(ax, GUARD_SENTENCE + " Inherited-host ZHAI omitted; censored "
                  "points omitted.", footnote)


FIGD_NBOOT = 2000                       # cluster-bootstrap draws (fixed seed)
FIGD_SEED = 42


def _figd_fit(points, n_boot=FIGD_NBOOT, seed=FIGD_SEED):
    """Combined log-log OLS of Smelt on apatite XF/XCl over the Fig D point
    set (matrix + magmatic-host ZHAI, non-censored, ratio present); ruling
    2026-08-06: ONE fit across both populations, declared in the caption.

    Returns None if fewer than 8 usable points. Uncertainty by CLUSTER
    bootstrap: sample_ids are resampled with replacement, never individual
    points -- all matrix points of a sample share that sample's k_med, so a
    plain point bootstrap would understate the CI (pseudo-replication). The
    envelope is the P16-P84 of residuals about the fit: an empirical
    prediction band, parallel in log space. log10(XF/XCl) is an ALR-type
    log-ratio coordinate of the closed X-site composition, so the regression
    is closure-safe; Smelt is not part of that closed system."""
    trip = [(p["xf_xcl"], p["smelt_med"], p["sample_id"])
            for typ in ("matrix", "ZHAI")
            for p in points if _figd_usable(p, typ)]
    if len(trip) < 8:
        return None
    lx = np.log10(np.array([t[0] for t in trip]))
    ly = np.log10(np.array([t[1] for t in trip]))
    sid = np.array([t[2] for t in trip])
    b, a = np.polyfit(lx, ly, 1)
    res = ly - (a + b * lx)
    e16, e84 = np.percentile(res, [16.0, 84.0])
    xg = np.linspace(lx.min(), lx.max(), 60)
    rng = np.random.default_rng(seed)
    uniq = np.unique(sid)
    slopes = np.empty(n_boot)
    lines = np.empty((n_boot, xg.size))
    kept = 0
    for _ in range(n_boot):
        pick = rng.choice(uniq, size=uniq.size, replace=True)
        idx = np.concatenate([np.flatnonzero(sid == s) for s in pick])
        if np.unique(lx[idx]).size < 2:
            continue
        bb, ab = np.polyfit(lx[idx], ly[idx], 1)
        slopes[kept] = bb
        lines[kept] = ab + bb * xg
        kept += 1
    lines = lines[:kept]
    ci_lo, ci_hi = np.percentile(lines, [2.5, 97.5], axis=0)
    b_lo, b_hi = np.percentile(slopes[:kept], [2.5, 97.5])
    return dict(a=a, b=b, b_ci=(float(b_lo), float(b_hi)), n=len(trip),
                n_samples=int(uniq.size), n_boot=kept, xg=xg,
                fit=a + b * xg, ci_lo=ci_lo, ci_hi=ci_hi,
                env_lo=a + b * xg + e16, env_hi=a + b * xg + e84,
                e16=float(e16), e84=float(e84))


def _figd_usable(p, typ):
    """Points that enter Fig D: non-censored, own XF/XCl present, and (ZHAI)
    magmatic host only -- inherited-host ZHAI are OMITTED per the 2026-08-06
    ruling (Fig C rule, not the Fig A grey-de-emphasis rule)."""
    return (p["type"] == typ and not p["censored"]
            and p.get("xf_xcl") is not None and p["xf_xcl"] > 0
            and p["smelt_med"] is not None
            and (typ != "ZHAI" or p["host_magmatic"]))


def draw_fig_d(ax, points, sample_k, footnote=True):
    """Fig D (manuscript candidate; rulings 2026-08-06, box/fit revision same
    day): Smelt vs the SAME analysis's apatite X-site mole ratio XF/XCl.

    x carries no input of the Smelt conversion (dFMQ/T enter y only), and in
    the P1 forward model XF/XCl rises along the degassing path (Cl partitions
    to the exsolved fluid while F stays in the melt), so the panel reads the
    S record against degassing progress recorded by the halogens. Layers:
    ONE combined log-log OLS line (see _figd_fit; slope 95% CI from the
    cluster bootstrap shown in the legend -- the CI/envelope BANDS were
    removed by the later 2026-08-06 ruling, the helper still computes them
    for the log diagnostics), analyses as small de-emphasized symbols, and
    per-unit-per-type VERTICAL BOXPLOTS of Smelt (Fig B styling, whis 5-95,
    n >= 3) positioned at the group's median XF/XCl (boxes replaced the v1
    median symbols per the same-day ruling). The caption MUST carry the
    guard sentence."""
    from matplotlib.lines import Line2D
    from matplotlib.patches import Patch

    fit = _figd_fit(points)
    if fit is not None:
        # bands REMOVED by user ruling 2026-08-06 (later same day): the line
        # alone is drawn; slope CI stays in the legend + run log, and
        # _figd_fit keeps computing CI/envelope for the log diagnostics.
        ax.plot(10.0 ** fit["xg"], 10.0 ** fit["fit"], color="#333333",
                linewidth=1.3, zorder=4)
        log(f"[figD] combined log-log fit: slope {fit['b']:.3f} "
            f"(95% CI {fit['b_ci'][0]:.3f}..{fit['b_ci'][1]:.3f}), "
            f"n={fit['n']} points / {fit['n_samples']} samples, "
            f"{fit['n_boot']} cluster-bootstrap draws; residual P16/P84 "
            f"{fit['e16']:+.3f}/{fit['e84']:+.3f} log10 units")
    n_noratio = 0
    for unit in UNIT_ORDER:
        for typ, marker, s_small in (("matrix", "o", 14), ("ZHAI", "^", 20)):
            pts = [p for p in points
                   if _figd_usable(p, typ) and p["unit"] == unit]
            n_noratio += sum(
                1 for p in points
                if p["type"] == typ and p["unit"] == unit
                and not p["censored"] and p["smelt_med"] is not None
                and (typ != "ZHAI" or p["host_magmatic"])
                and (p.get("xf_xcl") is None or p["xf_xcl"] <= 0))
            if pts:
                ax.scatter([p["xf_xcl"] for p in pts],
                           [p["smelt_med"] for p in pts], marker=marker,
                           s=s_small, alpha=0.45, zorder=2,
                           **_marker_kw(unit))
    for unit in UNIT_ORDER:                 # boxes on top (replace v1 medians)
        for typ in ("matrix", "ZHAI"):
            pts = [p for p in points
                   if _figd_usable(p, typ) and p["unit"] == unit]
            if len(pts) < 3:
                continue
            pos = float(np.median([p["xf_xcl"] for p in pts]))
            bp = ax.boxplot([[p["smelt_med"] for p in pts]], positions=[pos],
                            widths=[pos * 0.30], whis=(5, 95),
                            showfliers=False, patch_artist=True,
                            manage_ticks=False, zorder=5)
            bp["boxes"][0].set(facecolor=UNIT_COLOR[unit], alpha=0.18,
                               edgecolor="#666666", linewidth=0.8)
            for part in ("medians", "whiskers", "caps"):
                for art in bp[part]:
                    art.set(color="#666666", linewidth=0.8)
            bp["medians"][0].set(color="#222222", linewidth=1.2)
    ax.set_xscale("log")
    ax.set_yscale("log")
    ax.set_xlabel("apatite XF/XCl (X-site mole ratio)", fontsize=9)
    ax.set_ylabel("melt S (ppm)", fontsize=9)
    ax.tick_params(labelsize=7)
    ax.set_title(FIG_TITLES["D"], fontsize=9)
    handles = [Patch(facecolor=UNIT_COLOR[u], alpha=0.45,
                     edgecolor="#666666", label=UNIT_DISPLAY[u])
               for u in UNIT_ORDER]
    if fit is not None:
        handles.append(Line2D([], [], color="#333333", linewidth=1.3,
                              label=(f"combined fit, slope {fit['b']:.2f} "
                                     f"[{fit['b_ci'][0]:.2f}, "
                                     f"{fit['b_ci'][1]:.2f}] "
                                     f"(cluster bootstrap)")))
    ax.legend(handles=handles, fontsize=6.2, ncol=2, frameon=False,
              loc="lower left")
    note = ("One log-log OLS across BOTH populations (ruling 2026-08-06) -- "
            "the trend includes the between-population contrast (ZHAI k per "
            "host grain, matrix k per sample); slope 95% CI by cluster "
            "bootstrap resampling samples (not points, which share sample "
            "k). Boxes (whis 5-95%) sit at each group's median XF/XCl. "
            "XF/XCl rises along the modelled degassing path and shares no "
            "input with the conversion; inherited-host ZHAI and censored "
            "points omitted. " + GUARD_SENTENCE)
    if n_noratio:
        note += f" {n_noratio} point(s) lack XF/XCl and are not shown."
    _footnote(ax, note, footnote)


FIG_SIZES = {"A": (6.4, 4.6), "B": (8.2, 4.8), "C": (6.4, 4.6),
             "D": (6.4, 4.6)}


def write_figures(out_dir, points, sample_k, keys=FIG_KEYS):
    """CLI path: render each figure headless and save PDF + PNG."""
    import matplotlib
    matplotlib.use("Agg")
    import matplotlib.pyplot as plt
    for key in keys:
        fig, ax = plt.subplots(figsize=FIG_SIZES[key])
        draw_figure(key, ax, points, sample_k)
        _save_fig(fig, out_dir, FIG_STEMS[key])
        plt.close(fig)


# ===========================================================================
# ==== validation + golden ====
# ===========================================================================
def validate():
    """Spec section 8 spot checks; run at every startup, refuse on failure."""
    checks = []
    k = k_factor(1.5, 850.0)
    checks.append(("k(dFMQ=1.5, T=850C)", k, 0.4597, 5e-4))
    checks.append(("Smelt(S_Ap=1000, same)", 1000.0 * k, 460.0, 1.0))
    checks.append(("S_Ap(SO3=0.10 wt%)", sap_ppm(0.10), 400.5, 1e-6))
    ok = True
    for name, got, want, tol in checks:
        good = abs(got - want) <= tol
        ok = ok and good
        log(f"[self-check] {name}: got {got:.4f}, expect {want:g} "
              f"(tol {tol:g}) -> {'PASS' if good else 'FAIL'}")
    return ok


def gd_pair(d, t):
    return f"(dFMQ {d:.3f}, T {t:.1f} degC)"


def golden_compare(points):
    """Two-tier comparison against the reference Excel calculator.

    Tier 1 recomputes all 49 rows from the sheet's own S_Ap/dFMQ/T columns.
    Rows 11-12 carry a documented off-by-one in the sheet's H formula
    (H11 = C12*G11, H12 = C13*G12): the sheet value must match the WRONG
    formula and our recomputation the corrected one.
    Tier 2 aligns our ZHAI results by normalized inclusion id; the 5
    Edinburgh rows and the 3 unpaired 24OT02-14 rows (the sheet re-pairs
    them from an undocumented source; the spec says drop) are
    expected-absent. S_Ap tolerance 3e-4 covers the sheet's 4004.0 factor
    vs canonical 4005.0.
    Third documented divergence: for 24OT01-1-02-zir the sheet substituted
    the host values (-2.087, 693.2 degC) for the pairing table's actual
    host (-7.393, 1142.9 degC; a 1747 Ma inherited grain). The pairing
    table Zircon-apa-OT.xlsx is canonical (spec section 2.1: one-to-one
    pairing); our value follows it, the row is verified for arithmetic
    self-consistency and reported instead of compared.
    """
    if not os.path.exists(GOLDEN_XLSX):
        sys.exit(f"[FATAL] golden workbook missing: {GOLDEN_XLSX}")
    from openpyxl import load_workbook
    ws = load_workbook(GOLDEN_XLSX, data_only=True)["Calculator"]
    coef = [ws[f"K{i}"].value for i in range(2, 6)]
    if coef != [COEF_DFMQ, COEF_INTERCEPT, COEF_T, T_REF_C]:
        sys.exit(f"[FATAL] golden coefficients {coef} differ from module "
                 f"constants")
    rows = []
    for i in range(11, 60):
        vals = [ws.cell(row=i, column=c).value for c in range(2, 12)]
        if vals[0] is not None:
            rows.append((i, vals))
    if all(v[1][6] is None for v in rows):
        sys.exit("[FATAL] golden S_melt column is empty -- workbook was "
                 "saved without cached values; open+save it in Excel once.")
    OFF_BY_ONE = {11: 12, 12: 13}   # sheet row -> row whose C feeds its H
    n_ok = n_bug = 0
    for idx, v in rows:
        name, sap, dfmq, t, expo, dcorr, smelt = v[0], v[1], v[2], v[3], \
            v[4], v[5], v[6]
        my_k = k_factor(dfmq, t)
        assert abs(my_k - dcorr) <= 1e-9 * max(1.0, abs(dcorr)), \
            f"golden k mismatch at row {idx} ({name})"
        my_smelt = sap * my_k
        if idx in OFF_BY_ONE:
            wrong_sap = ws.cell(row=OFF_BY_ONE[idx], column=3).value
            assert abs(wrong_sap * dcorr - smelt) <= 1e-6 * max(1.0, smelt),\
                f"row {idx}: sheet value no longer matches the documented " \
                f"off-by-one bug -- workbook changed, re-audit"
            log(f"[golden] tier1 row {idx} ({name}): documented Excel "
                  f"off-by-one (sheet {smelt:.4g}, corrected "
                  f"{my_smelt:.4g})")
            n_bug += 1
            continue
        assert abs(my_smelt - smelt) <= 1e-9 * max(1.0, abs(smelt)), \
            f"golden Smelt mismatch at row {idx} ({name})"
        n_ok += 1
    log(f"[golden] tier1: {n_ok} rows exact + {n_bug} documented-bug rows "
          f"of {len(rows)}")

    ours = {normalize_id(p["point_id"]): p for p in points
            if p["type"] == "ZHAI"}
    # sheet hand-edits where the golden host differs from the canonical
    # pairing table (see docstring); arithmetic-checked, not compared.
    host_substituted = {"24OT01-1-2-zir"}
    n_cmp = n_cen = 0
    absent = []
    max_sap_rel = 0.0
    golden_zero = set()
    for idx, v in rows:
        name, sap, dfmq, t = v[0], v[1], v[2], v[3]
        key = normalize_id(name)
        p = ours.get(key)
        if p is None:
            absent.append(name)
            continue
        if not sap:
            golden_zero.add(key)
            assert p["censored"], f"{name}: golden S_Ap=0 but ours not " \
                                  f"censored"
            n_cen += 1
            continue
        rel = abs(p["s_ap_ppm"] - sap) / sap
        max_sap_rel = max(max_sap_rel, rel)
        assert rel <= 3e-4, f"{name}: S_Ap rel diff {rel:.2e} > 3e-4"
        if key in host_substituted:
            own = p["s_ap_ppm"] * k_factor(p["host_dfmq"], p["host_t_c"])
            assert abs(p["smelt_med"] - own) <= 1e-9 * own, \
                f"{name}: pairing-table arithmetic inconsistent"
            log(f"[golden] tier2 {name}: documented host substitution in "
                  f"the sheet (golden {gd_pair(dfmq, t)}, canonical pairing "
                  f"{gd_pair(p['host_dfmq'], p['host_t_c'])}) -- ours "
                  f"follows the pairing table")
            n_cmp += 1
            continue
        my_ref = sap * k_factor(dfmq, t) * (SO3_TO_S_PPM / 4004.0)
        rel_s = abs(p["smelt_med"] - my_ref) / my_ref
        assert rel_s <= 1e-9, f"{name}: Smelt rel diff {rel_s:.2e}"
        n_cmp += 1
    our_cen = {k for k, p in ours.items() if p["censored"]}
    assert our_cen == golden_zero, \
        f"censored set mismatch: ours-only {our_cen - golden_zero}, " \
        f"golden-only {golden_zero - our_cen}"
    log(f"[golden] tier2: {n_cmp} non-censored + {n_cen} censored aligned; "
          f"{len(absent)} expected-absent (Edinburgh/unpaired): "
          f"{', '.join(absent)}")
    log(f"[golden] tier2 max S_Ap rel diff {max_sap_rel:.2e} "
          f"(sheet factor 4004.0 vs canonical {SO3_TO_S_PPM:g})")
    log("[golden] PASS")


def _selfcheck_alteration():
    try:
        import apatite_m1
    except Exception as e:  # stripped env: keep running, loudly
        log(f"[WARNING] apatite_m1 unimportable ({e}); local "
              f"ALTERATION_SAMPLES mirror unverified")
        return
    if set(apatite_m1.ALTERATION_SAMPLES) != ALTERATION_SAMPLES:
        sys.exit("[FATAL] ALTERATION_SAMPLES mirror diverged from "
                 "apatite_m1 (single-source rule)")


# ===========================================================================
# ==== analysis driver (shared by the CLI and the GUI panel) ====
# ===========================================================================
def run_analysis(out_dir=None, reprep=False, write=True, sources=None):
    """Full pipeline -> dict(points, sample_k, unit_stats, meta, unpaired).

    sources: optional {tag: path} raw-input overrides, see run_prep.
    Callers embedding this (the GUI panel) must catch SystemExit: the data
    guards abort via sys.exit by design, which is fatal for a CLI but must
    become a message box in a window.
    """
    out_dir = out_dir or OUT_DIR
    os.makedirs(out_dir, exist_ok=True)
    paths = run_prep(out_dir, reprep, sources=sources)

    zircon = _read_input_csv(paths["zircon"])
    zhai = _read_input_csv(paths["zhai"])
    matrix = _read_input_csv(paths["matrix"])
    log(f"[input] normalized rows: zircon {len(zircon)}, zhai {len(zhai)},"
        f" matrix {len(matrix)}")
    guard_edinburgh({"zircon": zircon, "zhai": zhai, "matrix": matrix})
    _assert_unique_points(zhai, "inclusion_id", "zhai")
    _assert_unique_points(matrix, "point_id", "matrix")
    for label, rows in (("zhai", zhai), ("matrix", matrix)):
        if rows and "xf_xcl" not in rows[0]:
            sys.exit(f"[FATAL] cached {label} input predates the xf_xcl "
                     f"column (Fig D, 2026-08-06) -- re-run with --reprep")

    screened = screen_zircons(zircon)
    by_sample, by_unit = build_k_pools(screened)
    for u in UNIT_ORDER:
        log(f"[screen]   unit pool {u}: n={len(by_unit.get(u, []))}")

    zhai_pts, unpaired = compute_zhai_points(zhai)
    mx_pts, sample_k = compute_matrix_points(matrix, by_sample, by_unit)
    points = zhai_pts + mx_pts
    n_cen = sum(1 for p in points if p["censored"])
    log(f"[censored] {n_cen} point(s) with SO3 <= d.l. (no Smelt; see QC)")
    ustats = unit_stats(points)

    # Edinburgh prep counts for the QC report, read back from the normalized
    # inputs' provenance lines (parsing the log stream would be fragile)
    meta = {"ed_zircon": "-", "ed_zhai": "-", "ed_matrix": "-"}
    for tag in ("zircon", "zhai", "matrix"):
        with open(paths[tag], "r", encoding="utf-8") as f:
            for ln in f:
                if ln.startswith("# edinburgh excluded:"):
                    meta[f"ed_{tag}"] = ln.split(":")[1].strip()
                if not ln.startswith("#"):
                    break

    if write:
        write_outputs(out_dir, points, sample_k, ustats)
        write_qc(out_dir, meta, points, sample_k, unpaired)
    return {"points": points, "sample_k": sample_k, "unit_stats": ustats,
            "meta": meta, "unpaired": unpaired, "out_dir": out_dir}


# ===========================================================================
# ==== main ====
# ===========================================================================
def main(argv=None):
    try:
        sys.stdout.reconfigure(encoding="utf-8", errors="replace")
    except AttributeError:
        pass
    ap = argparse.ArgumentParser(
        description="Ok Tedi P1 apatite S -> melt S (Tassara et al. 2026)")
    ap.add_argument("--validate", action="store_true",
                    help="run spec section-8 spot checks and exit")
    ap.add_argument("--golden", action="store_true",
                    help="compare against the reference Excel calculator")
    ap.add_argument("--no-fig", action="store_true",
                    help="skip figure generation")
    ap.add_argument("--out", default=OUT_DIR, help="output directory")
    ap.add_argument("--prep-only", action="store_true",
                    help="build normalized input CSVs and exit")
    ap.add_argument("--reprep", action="store_true",
                    help="force rebuild of the normalized input CSVs")
    args = ap.parse_args(argv)

    if not validate():
        sys.exit("[FATAL] spot checks failed -- refusing to run")
    if args.validate:
        return
    _selfcheck_alteration()

    out_dir = args.out
    if args.prep_only:
        os.makedirs(out_dir, exist_ok=True)
        run_prep(out_dir, args.reprep)
        return

    res = run_analysis(out_dir, reprep=args.reprep)
    if args.golden:
        golden_compare(res["points"])
    if not args.no_fig:
        write_figures(out_dir, res["points"], res["sample_k"])


if __name__ == "__main__":
    main()
